from __future__ import annotations

import hashlib
import hmac
import io
import json
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Annotated

from fastapi import APIRouter, Body, Depends, Request
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from config.get_db import get_db
from module_admin.service.login_service import LoginService
from module_material.service.material_standardization_service import (
    STANDARD_STATUS_PENDING,
    create_review_task_for_price,
    enqueue_standardization,
    ensure_standardization_tables,
)
from utils.response_util import ResponseUtil

material_price_controller = APIRouter(
    prefix='/material/price', dependencies=[Depends(LoginService.get_current_user)]
)

TEMPLATE_HEADERS = [
    '*材料名称',
    '规格型号',
    '所属分类',
    '工艺段',
    '*单位',
    '品牌',
    '*数量',
    '不含税价',
    '税率(%)',
    '*含税价',
    '来源项目',
    '*采购时间(YYYY-MM)',
    '*价格类型',
    '报价供应商',
    '备注',
]

PRICE_TYPES = {'投标价', '中标价', '厂商报价', '市场询价', '信息价', '项目价'}
MONTH_REGEX = re.compile(r'^\d{4}-(0[1-9]|1[0-2])$')
FULL_DATE_REGEX = re.compile(r'^(\d{4}-(0[1-9]|1[0-2]))-\d{2}$')
SIGN_SECRET = 'material-price-import-sign-v1'


def _to_int(value: Any, default: int) -> int:
    try:
        if value is None:
            return default
        return int(value)
    except Exception:
        return default


def _to_decimal(value: Any, default: Decimal | None = None) -> Decimal | None:
    if value is None:
        return default
    s = str(value).strip()
    if not s:
        return default
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return None


def _as_str(value: Any) -> str:
    return str(value or '').strip()


def _normalize_purchase_month(value: Any) -> str:
    s = _as_str(value)
    if not s:
        return ''
    if MONTH_REGEX.match(s):
        return s
    full_date = FULL_DATE_REGEX.match(s)
    if full_date:
        return full_date.group(1)
    return s


def _round2(v: Decimal) -> Decimal:
    return v.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def _safe_json_loads(value: Any) -> Any:
    if value in (None, ''):
        return None
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(str(value))
    except Exception:
        return None


def _build_standard_spec_summary(attr_ids_json: Any, attr_value_map: dict[int, dict]) -> str:
    attr_ids = _safe_json_loads(attr_ids_json)
    if not isinstance(attr_ids, list):
        return ''
    parts: list[str] = []
    for item in attr_ids:
        try:
            attr_id = int(item)
        except Exception:
            continue
        attr_row = attr_value_map.get(attr_id)
        if not attr_row:
            continue
        spec_key = _as_str(attr_row.get('specKey'))
        spec_value = _as_str(attr_row.get('specValue'))
        if spec_key and spec_value:
            parts.append(f'{spec_key}:{spec_value}')
        elif spec_value:
            parts.append(spec_value)
    return '；'.join(parts)


def _to_float_or_none(value: Any) -> float | None:
    dec = _to_decimal(value)
    if dec is None:
        return None
    try:
        return float(dec)
    except Exception:
        return None


def _canonical_rows(rows: list[dict]) -> str:
    return json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def _sign_rows(rows: list[dict]) -> str:
    raw = _canonical_rows(rows).encode('utf-8')
    return hmac.new(SIGN_SECRET.encode('utf-8'), raw, hashlib.sha256).hexdigest()


async def _fetch_all(db: AsyncSession, sql: str, params: dict | None = None) -> list[dict]:
    res = await db.execute(text(sql), params or {})
    return [dict(row) for row in res.mappings().all()]


async def _fetch_scalar(db: AsyncSession, sql: str, params: dict | None = None) -> Any:
    res = await db.execute(text(sql), params or {})
    return res.scalar()


async def _column_exists(db: AsyncSession, table_name: str, column_name: str) -> bool:
    count = await _fetch_scalar(
        db,
        """
        SELECT COUNT(1)
        FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA = DATABASE()
          AND TABLE_NAME = :table_name
          AND COLUMN_NAME = :column_name
        """,
        {'table_name': table_name, 'column_name': column_name},
    )
    return int(count or 0) > 0


async def _ensure_column(db: AsyncSession, table_name: str, column_name: str, ddl: str) -> None:
    if await _column_exists(db, table_name, column_name):
        return
    await db.execute(text(f"ALTER TABLE {table_name} ADD COLUMN {ddl}"))


async def _ensure_tables(db: AsyncSession) -> None:
    await db.execute(
        text(
            """
            CREATE TABLE IF NOT EXISTS material_price_record (
                id BIGINT PRIMARY KEY AUTO_INCREMENT,
                material_name VARCHAR(255) NOT NULL,
                specification VARCHAR(255) NOT NULL,
                category_name VARCHAR(255) NOT NULL,
                category_level1_id BIGINT NULL,
                category_level2_id BIGINT NULL,
                category_level3_id BIGINT NULL,
                process_segment_name VARCHAR(255) NULL,
                unit_name VARCHAR(64) NOT NULL,
                brand_name VARCHAR(128) NULL,
                quantity DECIMAL(18, 4) NOT NULL DEFAULT 1,
                price_excluding_tax DECIMAL(18, 2) NOT NULL,
                tax_rate DECIMAL(6, 2) NOT NULL,
                price_including_tax DECIMAL(18, 2) NOT NULL,
                source_project VARCHAR(255) NULL,
                purchase_time VARCHAR(7) NOT NULL,
                price_type VARCHAR(20) NOT NULL,
                supplier_company VARCHAR(255) NULL,
                remark VARCHAR(500) NULL,
                batch_id VARCHAR(64) NULL,
                del_flag CHAR(1) NOT NULL DEFAULT '0',
                create_by VARCHAR(64) NULL,
                update_by VARCHAR(64) NULL,
                create_time DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                update_time DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
                KEY idx_material_price_key (material_name, specification, purchase_time),
                KEY idx_material_price_supplier (supplier_company),
                KEY idx_material_price_batch (batch_id)
            )
            """
        )
    )
    await _ensure_column(
        db,
        'material_price_record',
        'process_segment_name',
        "process_segment_name VARCHAR(255) NULL COMMENT '工艺段'",
    )
    await _ensure_column(
        db,
        'material_price_record',
        'category_level1_id',
        "category_level1_id BIGINT NULL COMMENT '分类一级ID'",
    )
    await _ensure_column(
        db,
        'material_price_record',
        'category_level2_id',
        "category_level2_id BIGINT NULL COMMENT '分类二级ID'",
    )
    await _ensure_column(
        db,
        'material_price_record',
        'category_level3_id',
        "category_level3_id BIGINT NULL COMMENT '分类三级ID'",
    )
    await ensure_standardization_tables(db)
    await db.execute(
        text(
            """
            CREATE TABLE IF NOT EXISTS material_price_category_dict (
                id BIGINT PRIMARY KEY AUTO_INCREMENT,
                category_name VARCHAR(255) NOT NULL,
                create_time DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE KEY uk_material_price_category_name (category_name)
            )
            """
        )
    )
    await db.execute(
        text(
            """
            CREATE TABLE IF NOT EXISTS material_price_unit_dict (
                id BIGINT PRIMARY KEY AUTO_INCREMENT,
                unit_name VARCHAR(64) NOT NULL,
                create_time DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE KEY uk_material_price_unit_name (unit_name)
            )
            """
        )
    )
    await db.execute(
        text(
            """
            CREATE TABLE IF NOT EXISTS material_price_supplier_dict (
                id BIGINT PRIMARY KEY AUTO_INCREMENT,
                supplier_name VARCHAR(255) NOT NULL,
                create_time DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE KEY uk_material_price_supplier_name (supplier_name)
            )
            """
        )
    )
    await db.execute(
        text(
            """
            CREATE TABLE IF NOT EXISTS material_price_brand_dict (
                id BIGINT PRIMARY KEY AUTO_INCREMENT,
                brand_name VARCHAR(255) NOT NULL,
                create_time DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE KEY uk_material_price_brand_name (brand_name)
            )
            """
        )
    )
    await db.execute(
        text(
            """
            CREATE TABLE IF NOT EXISTS material_price_project_dict (
                id BIGINT PRIMARY KEY AUTO_INCREMENT,
                project_name VARCHAR(255) NOT NULL,
                create_time DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE KEY uk_material_price_project_name (project_name)
            )
            """
        )
    )


async def _load_dict_sets(db: AsyncSession) -> tuple[set[str], set[str], set[str], set[str], set[str]]:
    category_rows = await _fetch_all(db, "SELECT category_name AS name FROM material_category WHERE status = '1'")
    category_rows += await _fetch_all(db, "SELECT category_name AS name FROM material_price_category_dict")
    categories = {_as_str(i['name']) for i in category_rows if _as_str(i.get('name'))}

    unit_rows = await _fetch_all(db, "SELECT unit_name AS name FROM material_unit WHERE status = '1'")
    unit_rows += await _fetch_all(db, "SELECT unit_name AS name FROM material_price_unit_dict")
    units = {_as_str(i['name']) for i in unit_rows if _as_str(i.get('name'))}

    suppliers = {
        _as_str(i['name'])
        for i in await _fetch_all(
            db,
            """
            SELECT supplier_name AS name FROM material_price_supplier_dict
            """,
        )
        if _as_str(i.get('name'))
    }
    suppliers |= {
        _as_str(i['name'])
        for i in await _fetch_all(db, "SELECT DISTINCT supplier_company AS name FROM material_price_record WHERE COALESCE(supplier_company, '') <> ''")
        if _as_str(i.get('name'))
    }

    brands = {
        _as_str(i['name'])
        for i in await _fetch_all(
            db,
            """
            SELECT brand_name AS name FROM material_price_brand_dict
            """,
        )
        if _as_str(i.get('name'))
    }
    brands |= {
        _as_str(i['name'])
        for i in await _fetch_all(db, "SELECT DISTINCT brand_name AS name FROM material_price_record WHERE COALESCE(brand_name, '') <> ''")
        if _as_str(i.get('name'))
    }

    projects = {
        _as_str(i['name'])
        for i in await _fetch_all(
            db,
            """
            SELECT project_name AS name FROM material_price_project_dict
            """,
        )
        if _as_str(i.get('name'))
    }
    projects |= {
        _as_str(i['name'])
        for i in await _fetch_all(db, "SELECT DISTINCT source_project AS name FROM material_price_record WHERE COALESCE(source_project, '') <> ''")
        if _as_str(i.get('name'))
    }
    return categories, units, suppliers, brands, projects


async def _resolve_category_by_level3_id(db: AsyncSession, level3_id: int) -> dict | None:
    rows = await _fetch_all(
        db,
        """
        SELECT c3.id AS categoryLevel3Id,
               c3.category_name AS categoryName,
               c2.id AS categoryLevel2Id,
               c1.id AS categoryLevel1Id
        FROM material_category c3
        LEFT JOIN material_category c2 ON c2.id = c3.parent_id
        LEFT JOIN material_category c1 ON c1.id = c2.parent_id
        WHERE c3.id = :id AND c3.status = '1'
        LIMIT 1
        """,
        {'id': level3_id},
    )
    return rows[0] if rows else None


async def _resolve_category_by_name(db: AsyncSession, category_name: str) -> dict | None:
    rows = await _fetch_all(
        db,
        """
        SELECT c3.id AS categoryLevel3Id,
               c3.category_name AS categoryName,
               c2.id AS categoryLevel2Id,
               c1.id AS categoryLevel1Id
        FROM material_category c3
        LEFT JOIN material_category c2 ON c2.id = c3.parent_id
        LEFT JOIN material_category c1 ON c1.id = c2.parent_id
        WHERE c3.category_name = :name
          AND c3.status = '1'
          AND c3.level = 3
        ORDER BY c3.id ASC
        LIMIT 1
        """,
        {'name': category_name},
    )
    return rows[0] if rows else None


async def _exists_duplicate_in_db(db: AsyncSession, row: dict, exclude_id: int | None = None) -> bool:
    where_exclude = ''
    params = {
        'material_name': row['materialName'],
        'specification': row['specification'],
        'purchase_time': row['purchaseTime'],
        'source_project': _as_str(row.get('sourceProject')),
        'supplier_company': _as_str(row.get('supplierCompany')),
        'price_type': row['priceType'],
    }
    if exclude_id:
        where_exclude = ' AND id <> :exclude_id'
        params['exclude_id'] = exclude_id

    count = await _fetch_scalar(
        db,
        f"""
        SELECT COUNT(1)
        FROM material_price_record
        WHERE del_flag = '0'
          AND material_name = :material_name
          AND specification = :specification
          AND purchase_time = :purchase_time
          AND COALESCE(source_project, '') = :source_project
          AND COALESCE(supplier_company, '') = :supplier_company
          AND price_type = :price_type
          {where_exclude}
        """,
        params,
    )
    return (count or 0) > 0


def _normalize_row(raw: dict) -> dict:
    return {
        'materialName': _as_str(raw.get('materialName')),
        'specification': _as_str(raw.get('specification')),
        'categoryName': _as_str(raw.get('categoryName')),
        'categoryLevel1Id': _to_int(raw.get('categoryLevel1Id'), 0),
        'categoryLevel2Id': _to_int(raw.get('categoryLevel2Id'), 0),
        'categoryLevel3Id': _to_int(raw.get('categoryLevel3Id'), 0),
        'processSegmentName': _as_str(raw.get('processSegmentName')),
        'unit': _as_str(raw.get('unit')),
        'brand': _as_str(raw.get('brand')),
        'quantity': _as_str(raw.get('quantity')),
        'priceExcludingTax': _as_str(raw.get('priceExcludingTax')),
        'taxRate': _as_str(raw.get('taxRate')),
        'priceIncludingTax': _as_str(raw.get('priceIncludingTax')),
        'sourceProject': _as_str(raw.get('sourceProject')),
        'purchaseTime': _normalize_purchase_month(raw.get('purchaseTime')),
        'priceType': _as_str(raw.get('priceType')),
        'supplierCompany': _as_str(raw.get('supplierCompany')),
        'remark': _as_str(raw.get('remark')),
    }


def _key_compare_text(value: Any) -> str:
    return _as_str(value)


def _key_compare_int(value: Any) -> int:
    return _to_int(value, 0)


def _get_restandardize_changed_fields(current_row: dict, normalized_row: dict) -> list[str]:
    checks = [
        ('material_name', 'materialName', '材料名称', _key_compare_text),
        ('specification', 'specification', '规格型号', _key_compare_text),
        ('unit_name', 'unit', '单位', _key_compare_text),
        ('category_level3_id', 'categoryLevel3Id', '三级分类', _key_compare_int),
        ('process_segment_name', 'processSegmentName', '工艺段', _key_compare_text),
        ('remark', 'remark', '备注', _key_compare_text),
    ]
    changed: list[str] = []
    for old_key, new_key, label, cast_fn in checks:
        old_value = cast_fn(current_row.get(old_key))
        new_value = cast_fn(normalized_row.get(new_key))
        if old_value != new_value:
            changed.append(label)
    return changed


async def _validate_rows(db: AsyncSession, rows: list[dict], with_db_duplicate: bool = True) -> dict:
    normalized_rows: list[dict] = []
    result_rows: list[dict] = []

    for idx, raw in enumerate(rows, start=1):
        normalized = _normalize_row(raw)
        errors: list[dict] = []
        warnings: list[dict] = []

        resolved_category: dict | None = None
        if normalized['categoryLevel3Id'] > 0:
            resolved_category = await _resolve_category_by_level3_id(db, normalized['categoryLevel3Id'])
            if not resolved_category:
                errors.append({'field': 'categoryLevel3Id', 'message': '所属分类无效，请重新选择三级分类'})
        elif normalized['categoryName']:
            resolved_category = await _resolve_category_by_name(db, normalized['categoryName'])
            if not resolved_category:
                warnings.append({'field': 'categoryName', 'message': '未匹配到标准三级分类，将按原文本入库'})

        if resolved_category:
            normalized['categoryName'] = _as_str(resolved_category.get('categoryName'))
            normalized['categoryLevel1Id'] = _to_int(resolved_category.get('categoryLevel1Id'), 0)
            normalized['categoryLevel2Id'] = _to_int(resolved_category.get('categoryLevel2Id'), 0)
            normalized['categoryLevel3Id'] = _to_int(resolved_category.get('categoryLevel3Id'), 0)

        # 必填
        for field, label in [
            ('materialName', '材料名称'),
            ('unit', '单位'),
            ('quantity', '数量'),
            ('priceIncludingTax', '含税价'),
            ('purchaseTime', '采购时间'),
            ('priceType', '价格类型'),
        ]:
            if not normalized[field]:
                errors.append({'field': field, 'message': f'{label}不能为空'})

        # 数值
        quantity = _to_decimal(normalized['quantity'])
        if quantity is None:
            errors.append({'field': 'quantity', 'message': '数量必须是数字'})
        elif quantity <= Decimal('0'):
            errors.append({'field': 'quantity', 'message': '数量必须大于0'})

        excl = _to_decimal(normalized['priceExcludingTax'])
        incl = _to_decimal(normalized['priceIncludingTax'])
        tax = _to_decimal(normalized['taxRate'])

        if normalized['priceExcludingTax'] and excl is None:
            errors.append({'field': 'priceExcludingTax', 'message': '不含税价必须是数字'})
        elif excl is not None and excl < 0:
            errors.append({'field': 'priceExcludingTax', 'message': '不含税价不能小于0'})

        if incl is None:
            errors.append({'field': 'priceIncludingTax', 'message': '含税价必须是数字'})
        elif incl < 0:
            errors.append({'field': 'priceIncludingTax', 'message': '含税价不能小于0'})

        if normalized['taxRate'] and tax is None:
            errors.append({'field': 'taxRate', 'message': '税率必须是数字'})
        elif tax is not None and (tax < 0 or tax > 100):
            errors.append({'field': 'taxRate', 'message': '税率必须在0-100之间'})

        if excl is not None and incl is not None and tax is not None and excl >= 0 and incl >= 0 and 0 <= tax <= 100:
            expected = _round2(excl * (Decimal('1') + tax / Decimal('100')))
            if (expected - _round2(incl)).copy_abs() > Decimal('0.05'):
                errors.append({'field': 'priceIncludingTax', 'message': f'含税价与不含税价/税率不一致，应约为{expected}'})

        if normalized['purchaseTime'] and not MONTH_REGEX.match(normalized['purchaseTime']):
            errors.append({'field': 'purchaseTime', 'message': '采购时间格式必须为YYYY-MM'})

        if normalized['priceType'] and normalized['priceType'] not in PRICE_TYPES:
            errors.append({'field': 'priceType', 'message': '价格类型仅支持投标价/中标价/厂商报价/市场询价/信息价/项目价'})

        normalized_rows.append(normalized)
        result_rows.append(
            {
                'rowNo': idx,
                'rawData': raw,
                'normalizedData': normalized,
                'errors': errors,
                'warnings': warnings,
                'status': 'PASS' if not errors else 'FAIL',
            }
        )

    # 库内重复检查
    if with_db_duplicate:
        for item in result_rows:
            if item['status'] == 'FAIL':
                continue
            if await _exists_duplicate_in_db(db, item['normalizedData']):
                item['errors'].append({'field': 'materialName', 'message': '数据库已存在同一业务键记录'})
                item['status'] = 'FAIL'

    success_count = len([i for i in result_rows if i['status'] == 'PASS'])
    failed_count = len(result_rows) - success_count
    pass_rows = [i['normalizedData'] for i in result_rows if i['status'] == 'PASS']

    return {
        'summary': {
            'total': len(result_rows),
            'success': success_count,
            'failed': failed_count,
        },
        'rows': result_rows,
        'canCommit': failed_count == 0,
        'batchToken': _sign_rows(pass_rows),
    }


async def _upsert_import_dicts(db: AsyncSession, rows: list[dict]) -> None:
    for row in rows:
        category = _as_str(row.get('categoryName'))
        unit = _as_str(row.get('unit'))
        supplier = _as_str(row.get('supplierCompany'))
        brand = _as_str(row.get('brand'))
        project = _as_str(row.get('sourceProject'))

        if category:
            await db.execute(
                text(
                    """
                    INSERT INTO material_price_category_dict(category_name)
                    VALUES(:name)
                    ON DUPLICATE KEY UPDATE category_name = VALUES(category_name)
                    """
                ),
                {'name': category},
            )

        if unit:
            await db.execute(
                text(
                    """
                    INSERT INTO material_price_unit_dict(unit_name)
                    VALUES(:name)
                    ON DUPLICATE KEY UPDATE unit_name = VALUES(unit_name)
                    """
                ),
                {'name': unit},
            )

        if supplier:
            await db.execute(
                text(
                    """
                    INSERT INTO material_price_supplier_dict(supplier_name)
                    VALUES(:name)
                    ON DUPLICATE KEY UPDATE supplier_name = VALUES(supplier_name)
                    """
                ),
                {'name': supplier},
            )

        if brand:
            await db.execute(
                text(
                    """
                    INSERT INTO material_price_brand_dict(brand_name)
                    VALUES(:name)
                    ON DUPLICATE KEY UPDATE brand_name = VALUES(brand_name)
                    """
                ),
                {'name': brand},
            )

        if project:
            await db.execute(
                text(
                    """
                    INSERT INTO material_price_project_dict(project_name)
                    VALUES(:name)
                    ON DUPLICATE KEY UPDATE project_name = VALUES(project_name)
                    """
                ),
                {'name': project},
            )


@material_price_controller.get('/import/template')
@material_price_controller.post('/import/template')
async def material_price_import_template(
    request: Request,
    db: Annotated[AsyncSession, Depends(get_db)],
):
    await _ensure_tables(db)
    wb = Workbook()
    ws = wb.active
    ws.title = '材价导入模板'

    ws.append(TEMPLATE_HEADERS)
    ws.append(
        [
            '提升水泵',
            'Q=2083m3/h,H=9.5m,N=75kW',
            '通用设备',
            '回流污泥泵站',
            '台',
            'ABB',
            '6',
            '128000',
            '13',
            '144640',
            '市政管网改造工程',
            '2026-03',
            '投标价',
            '杭州某泵技术有限公司',
            '模板示例行，可删除',
        ]
    )

    widths = [18, 28, 14, 16, 8, 12, 10, 12, 10, 12, 24, 18, 10, 20, 20]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"material_price_import_template_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
    return ResponseUtil.streaming(
        data=stream,
        headers=headers,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@material_price_controller.post('/import/validate')
async def material_price_import_validate(
    request: Request,
    payload: Annotated[dict, Body()],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    await _ensure_tables(db)
    rows = payload.get('rows') or []
    if not isinstance(rows, list):
        return ResponseUtil.failure(msg='rows必须为数组')

    result = await _validate_rows(db, rows, with_db_duplicate=True)
    return ResponseUtil.success(data=result)


async def _list_master_names(db: AsyncSession, kind: str, keyword: str = '') -> list[str]:
    if kind == 'supplier':
        rows = await _fetch_all(
            db,
            """
            SELECT supplier_name AS name FROM material_price_supplier_dict
            UNION
            SELECT supplier_company AS name FROM material_price_record
            WHERE COALESCE(supplier_company, '') <> ''
            ORDER BY name ASC
            """,
        )
    elif kind == 'brand':
        rows = await _fetch_all(
            db,
            """
            SELECT brand_name AS name FROM material_price_brand_dict
            UNION
            SELECT brand_name AS name FROM material_price_record
            WHERE COALESCE(brand_name, '') <> ''
            ORDER BY name ASC
            """,
        )
    elif kind == 'project':
        rows = await _fetch_all(
            db,
            """
            SELECT project_name AS name FROM material_price_project_dict
            UNION
            SELECT source_project AS name FROM material_price_record
            WHERE COALESCE(source_project, '') <> ''
            ORDER BY name ASC
            """,
        )
    else:
        return []

    kw = keyword.lower().strip()
    names = []
    for item in rows:
        name = _as_str(item.get('name'))
        if not name:
            continue
        if kw and kw not in name.lower():
            continue
        names.append(name)
    return names[:200]


@material_price_controller.get('/master/supplier/list')
@material_price_controller.post('/master/supplier/list')
async def material_price_master_supplier_list(
    request: Request,
    db: Annotated[AsyncSession, Depends(get_db)],
    payload: Annotated[dict | None, Body()] = None,
):
    await _ensure_tables(db)
    keyword = _as_str((payload or {}).get('keyword') if isinstance(payload, dict) else request.query_params.get('keyword'))
    data = await _list_master_names(db, 'supplier', keyword)
    return ResponseUtil.success(data=data)


@material_price_controller.get('/master/brand/list')
@material_price_controller.post('/master/brand/list')
async def material_price_master_brand_list(
    request: Request,
    db: Annotated[AsyncSession, Depends(get_db)],
    payload: Annotated[dict | None, Body()] = None,
):
    await _ensure_tables(db)
    keyword = _as_str((payload or {}).get('keyword') if isinstance(payload, dict) else request.query_params.get('keyword'))
    data = await _list_master_names(db, 'brand', keyword)
    return ResponseUtil.success(data=data)


@material_price_controller.get('/master/project/list')
@material_price_controller.post('/master/project/list')
async def material_price_master_project_list(
    request: Request,
    db: Annotated[AsyncSession, Depends(get_db)],
    payload: Annotated[dict | None, Body()] = None,
):
    await _ensure_tables(db)
    keyword = _as_str((payload or {}).get('keyword') if isinstance(payload, dict) else request.query_params.get('keyword'))
    data = await _list_master_names(db, 'project', keyword)
    return ResponseUtil.success(data=data)


@material_price_controller.post('/import/commit')
async def material_price_import_commit(
    request: Request,
    payload: Annotated[dict, Body()],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    await _ensure_tables(db)
    rows = payload.get('rows') or []
    batch_token = _as_str(payload.get('batchToken'))
    if not isinstance(rows, list) or not rows:
        return ResponseUtil.failure(msg='提交数据不能为空')

    validate_result = await _validate_rows(db, rows, with_db_duplicate=True)
    if not validate_result['canCommit']:
        return ResponseUtil.failure(msg='存在未通过校验数据，禁止提交', data=validate_result)

    pass_rows = [i['normalizedData'] for i in validate_result['rows'] if i['status'] == 'PASS']
    expected_token = _sign_rows(pass_rows)
    if not batch_token or not hmac.compare_digest(batch_token, expected_token):
        return ResponseUtil.failure(msg='batchToken校验失败，请重新校验后提交')

    batch_id = hashlib.md5(f"{datetime.now().isoformat()}|{expected_token}".encode('utf-8')).hexdigest()
    create_by = _as_str(payload.get('createBy')) or 'system'
    inserted_ids: list[int] = []

    await _upsert_import_dicts(db, pass_rows)

    for row in pass_rows:
        await db.execute(
            text(
                """
                INSERT INTO material_price_record(
                    material_name, specification, category_name, category_level1_id, category_level2_id, category_level3_id, unit_name, brand_name,
                    process_segment_name,
                    quantity, price_excluding_tax, tax_rate, price_including_tax,
                    source_project, purchase_time, price_type, supplier_company,
                    remark, batch_id, del_flag, create_by, update_by, standardization_status
                ) VALUES(
                    :material_name, :specification, :category_name, :category_level1_id, :category_level2_id, :category_level3_id, :unit_name, :brand_name,
                    :process_segment_name,
                    :quantity, :price_excluding_tax, :tax_rate, :price_including_tax,
                    :source_project, :purchase_time, :price_type, :supplier_company,
                    :remark, :batch_id, '0', :create_by, :update_by, :standardization_status
                )
                """
            ),
            {
                'material_name': row['materialName'],
                'specification': row['specification'],
                'category_name': row['categoryName'],
                'category_level1_id': row.get('categoryLevel1Id') or None,
                'category_level2_id': row.get('categoryLevel2Id') or None,
                'category_level3_id': row.get('categoryLevel3Id') or None,
                'process_segment_name': row['processSegmentName'] or None,
                'unit_name': row['unit'],
                'brand_name': row['brand'] or None,
                'quantity': str(_to_decimal(row['quantity'], Decimal('0'))),
                'price_excluding_tax': str(_round2(_to_decimal(row['priceExcludingTax'], Decimal('0')))),
                'tax_rate': str(_round2(_to_decimal(row['taxRate'], Decimal('0')))),
                'price_including_tax': str(_round2(_to_decimal(row['priceIncludingTax'], Decimal('0')))),
                'source_project': row['sourceProject'] or None,
                'purchase_time': row['purchaseTime'],
                'price_type': row['priceType'],
                'supplier_company': row['supplierCompany'] or None,
                'remark': row.get('remark') or None,
                'batch_id': batch_id,
                'create_by': create_by,
                'update_by': create_by,
                'standardization_status': STANDARD_STATUS_PENDING,
            },
        )
        inserted_ids.append(int(await _fetch_scalar(db, "SELECT LAST_INSERT_ID()") or 0))

    for price_record_id in inserted_ids:
        await create_review_task_for_price(db, price_record_id)

    await db.commit()
    enqueue_standardization(inserted_ids)
    return ResponseUtil.success(data={'insertCount': len(pass_rows), 'batchId': batch_id})


@material_price_controller.post('/page')
async def material_price_page(
    request: Request,
    payload: Annotated[dict, Body()],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    await _ensure_tables(db)

    page_num = _to_int(payload.get('pageNum'), 1)
    page_size = _to_int(payload.get('pageSize'), 20)
    data_view_mode = _as_str(payload.get('dataViewMode')).upper() or 'NON_STANDARD'
    if data_view_mode not in {'NON_STANDARD', 'STANDARD'}:
        data_view_mode = 'NON_STANDARD'

    standard_mode = data_view_mode == 'STANDARD'
    field_prefix = 'p.' if standard_mode else ''

    def col(name: str) -> str:
        return f'{field_prefix}{name}'

    where = [f"{col('del_flag')} = '0'"]
    params: dict[str, Any] = {}

    def like_if(v: Any, field: str, key: str) -> None:
        s = _as_str(v)
        if s:
            where.append(f"{field} LIKE :{key}")
            params[key] = f'%{s}%'

    def eq_if(v: Any, field: str, key: str) -> None:
        s = _as_str(v)
        if s:
            where.append(f"{field} = :{key}")
            params[key] = s

    like_if(payload.get('materialName'), col('material_name'), 'material_name')
    like_if(payload.get('specification'), col('specification'), 'specification')
    like_if(payload.get('categoryName'), col('category_name'), 'category_name')
    eq_if(payload.get('brand'), col('brand_name'), 'brand_name')
    eq_if(payload.get('processSegmentName'), col('process_segment_name'), 'process_segment_name')
    eq_if(payload.get('unit'), col('unit_name'), 'unit_name')
    eq_if(payload.get('priceType'), col('price_type'), 'price_type')
    eq_if(payload.get('purchaseTime'), col('purchase_time'), 'purchase_time')
    eq_if(payload.get('standardCode'), col('standard_code'), 'standard_code')
    like_if(payload.get('sourceProject'), col('source_project'), 'source_project')
    like_if(payload.get('supplierCompany'), col('supplier_company'), 'supplier_company')
    if standard_mode:
        where.append(f"{col('standardization_status')} IN ('SUCCESS', 'MANUAL', 'SKIPPED')")

    if payload.get('taxRate') not in (None, '', 'null'):
        tax = _to_decimal(payload.get('taxRate'))
        if tax is not None:
            where.append(f"{col('tax_rate')} = :tax_rate")
            params['tax_rate'] = str(_round2(tax))

    where_sql = ' AND '.join(where)

    total_from = 'material_price_record p' if standard_mode else 'material_price_record'
    total = await _fetch_scalar(db, f"SELECT COUNT(1) FROM {total_from} WHERE {where_sql}", params) or 0

    params.update({'offset': (page_num - 1) * page_size, 'limit': page_size})

    if data_view_mode == 'STANDARD':
        rows = await _fetch_all(
            db,
            f"""
            SELECT
                p.id,
                p.material_name AS materialName,
                p.specification,
                p.category_name AS categoryName,
                p.category_level1_id AS categoryLevel1Id,
                p.category_level2_id AS categoryLevel2Id,
                p.category_level3_id AS categoryLevel3Id,
                p.process_segment_name AS processSegmentName,
                p.unit_name AS unit,
                p.brand_name AS brand,
                p.quantity,
                p.price_excluding_tax AS priceExcludingTax,
                p.tax_rate AS taxRate,
                p.price_including_tax AS priceIncludingTax,
                p.source_project AS sourceProject,
                p.purchase_time AS purchaseTime,
                p.price_type AS priceType,
                p.supplier_company AS supplierCompany,
                p.remark,
                p.batch_id AS batchId,
                p.create_time AS createTime,
                p.standardization_status AS standardizationStatus,
                p.standard_review_id AS standardReviewId,
                p.standard_code AS standardCode,
                p.standard_category_level3_id AS standardCategoryLevel3Id,
                p.standard_std_name_id AS standardStdNameId,
                p.standard_unit_id AS standardUnitId,
                p.standard_segment_id AS standardSegmentId,
                p.standard_attr_value_ids_json AS standardAttrValueIdsJson,
                p.normalized_price_excluding_tax AS normalizedPriceExcludingTax,
                p.normalized_price_including_tax AS normalizedPriceIncludingTax,
                p.standardization_error AS standardizationError,
                p.standardized_at AS standardizedAt,
                c3.category_name AS standardCategoryLevel3Name,
                sn.standard_name AS standardStdName,
                su.unit_name AS standardUnitName,
                ps.segment_name AS standardSegmentName
            FROM material_price_record p
            LEFT JOIN material_category c3 ON c3.id = p.standard_category_level3_id AND c3.status = '1'
            LEFT JOIN material_std_name sn ON sn.id = p.standard_std_name_id AND sn.del_flag = '0'
            LEFT JOIN material_unit su ON su.id = p.standard_unit_id AND su.status = '1'
            LEFT JOIN material_process_segment ps ON ps.id = p.standard_segment_id AND ps.status = '1'
            WHERE {where_sql}
            ORDER BY p.id DESC
            LIMIT :offset, :limit
            """,
            params,
        )
        attr_id_set: set[int] = set()
        for row in rows:
            attr_ids = _safe_json_loads(row.get('standardAttrValueIdsJson'))
            if isinstance(attr_ids, list):
                for item in attr_ids:
                    try:
                        attr_id = int(item)
                    except Exception:
                        continue
                    if attr_id > 0:
                        attr_id_set.add(attr_id)

        attr_value_map: dict[int, dict] = {}
        if attr_id_set:
            attr_params: dict[str, Any] = {}
            placeholders: list[str] = []
            for idx, attr_id in enumerate(sorted(attr_id_set)):
                key = f'attr_{idx}'
                placeholders.append(f':{key}')
                attr_params[key] = attr_id
            attr_rows = await _fetch_all(
                db,
                f"""
                SELECT
                    v.id,
                    v.spec_value AS specValue,
                    s.spec_key AS specKey
                FROM material_std_name_spec_value v
                LEFT JOIN material_std_name_spec s ON s.id = v.spec_id
                WHERE v.del_flag = '0'
                  AND v.id IN ({','.join(placeholders)})
                """,
                attr_params,
            )
            attr_value_map = {_to_int(item.get('id'), 0): item for item in attr_rows if _to_int(item.get('id'), 0) > 0}

        for row in rows:
            row['standardSpecSummary'] = _build_standard_spec_summary(row.get('standardAttrValueIdsJson'), attr_value_map)

        # 标准化视图下：按标准编码计算价格区间，且排除当前记录自身
        standard_codes = {
            _as_str(row.get('standardCode'))
            for row in rows
            if _as_str(row.get('standardCode'))
        }
        if standard_codes:
            code_params: dict[str, Any] = {}
            placeholders: list[str] = []
            for idx, code in enumerate(sorted(standard_codes)):
                key = f'code_{idx}'
                placeholders.append(f':{key}')
                code_params[key] = code

            peer_rows = await _fetch_all(
                db,
                f"""
                SELECT
                    id,
                    standard_code AS standardCode,
                    COALESCE(normalized_price_excluding_tax, price_excluding_tax) AS priceExcludingTax,
                    COALESCE(normalized_price_including_tax, price_including_tax) AS priceIncludingTax
                FROM material_price_record
                WHERE del_flag = '0'
                  AND standardization_status IN ('SUCCESS', 'MANUAL', 'SKIPPED')
                  AND COALESCE(standard_code, '') <> ''
                  AND standard_code IN ({','.join(placeholders)})
                """,
                code_params,
            )

            peers_by_code: dict[str, list[dict]] = {}
            for peer in peer_rows:
                code = _as_str(peer.get('standardCode'))
                if not code:
                    continue
                peers_by_code.setdefault(code, []).append(peer)

            for row in rows:
                row['rangeExcl'] = None
                row['rangeIncl'] = None

                code = _as_str(row.get('standardCode'))
                row_id = _to_int(row.get('id'), 0)
                if not code or row_id <= 0:
                    continue

                peers = peers_by_code.get(code, [])
                excl_values = [
                    price
                    for peer in peers
                    if _to_int(peer.get('id'), 0) != row_id
                    for price in [_to_float_or_none(peer.get('priceExcludingTax'))]
                    if price is not None
                ]
                incl_values = [
                    price
                    for peer in peers
                    if _to_int(peer.get('id'), 0) != row_id
                    for price in [_to_float_or_none(peer.get('priceIncludingTax'))]
                    if price is not None
                ]

                if excl_values:
                    row['rangeExcl'] = [min(excl_values), max(excl_values)]
                if incl_values:
                    row['rangeIncl'] = [min(incl_values), max(incl_values)]
    else:
        rows = await _fetch_all(
            db,
            f"""
            SELECT
                id,
                material_name AS materialName,
                specification,
                category_name AS categoryName,
                category_level1_id AS categoryLevel1Id,
                category_level2_id AS categoryLevel2Id,
                category_level3_id AS categoryLevel3Id,
                process_segment_name AS processSegmentName,
                unit_name AS unit,
                brand_name AS brand,
                quantity,
                price_excluding_tax AS priceExcludingTax,
                tax_rate AS taxRate,
                price_including_tax AS priceIncludingTax,
                source_project AS sourceProject,
                purchase_time AS purchaseTime,
                price_type AS priceType,
                supplier_company AS supplierCompany,
                remark,
                batch_id AS batchId,
                create_time AS createTime,
                standardization_status AS standardizationStatus,
                standard_review_id AS standardReviewId,
                standard_code AS standardCode,
                standard_category_level3_id AS standardCategoryLevel3Id,
                standard_std_name_id AS standardStdNameId,
                standard_unit_id AS standardUnitId,
                standard_segment_id AS standardSegmentId,
                standard_attr_value_ids_json AS standardAttrValueIdsJson,
                normalized_price_excluding_tax AS normalizedPriceExcludingTax,
                normalized_price_including_tax AS normalizedPriceIncludingTax,
                standardization_error AS standardizationError,
                standardized_at AS standardizedAt,
                '' AS standardCategoryLevel3Name,
                '' AS standardStdName,
                '' AS standardUnitName,
                '' AS standardSegmentName,
                '' AS standardSpecSummary
            FROM material_price_record
            WHERE {where_sql}
            ORDER BY id DESC
            LIMIT :offset, :limit
            """,
            params,
        )

    return ResponseUtil.success(data={'records': rows, 'total': int(total)})


@material_price_controller.post('/getById')
async def material_price_get_by_id(
    request: Request,
    payload: Annotated[dict, Body()],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    await _ensure_tables(db)
    item_id = _to_int(payload.get('id'), 0)
    if item_id <= 0:
        return ResponseUtil.failure(msg='ID不能为空')

    row = await _fetch_all(
        db,
        """
        SELECT
            id,
            material_name AS materialName,
            specification,
            category_name AS categoryName,
            category_level1_id AS categoryLevel1Id,
            category_level2_id AS categoryLevel2Id,
            category_level3_id AS categoryLevel3Id,
            process_segment_name AS processSegmentName,
            unit_name AS unit,
            brand_name AS brand,
            quantity,
            price_excluding_tax AS priceExcludingTax,
            tax_rate AS taxRate,
            price_including_tax AS priceIncludingTax,
            source_project AS sourceProject,
            purchase_time AS purchaseTime,
            price_type AS priceType,
            supplier_company AS supplierCompany,
            remark,
            standardization_status AS standardizationStatus,
            standard_review_id AS standardReviewId,
            standard_code AS standardCode,
            standard_category_level3_id AS standardCategoryLevel3Id,
            standard_std_name_id AS standardStdNameId,
            standard_unit_id AS standardUnitId,
            standard_segment_id AS standardSegmentId,
            standard_attr_value_ids_json AS standardAttrValueIdsJson,
            normalized_price_excluding_tax AS normalizedPriceExcludingTax,
            normalized_price_including_tax AS normalizedPriceIncludingTax,
            standardization_error AS standardizationError,
            standardized_at AS standardizedAt
        FROM material_price_record
        WHERE id = :id AND del_flag = '0'
        LIMIT 1
        """,
        {'id': item_id},
    )
    if not row:
        return ResponseUtil.failure(msg='材价记录不存在')
    return ResponseUtil.success(data=row[0])


@material_price_controller.post('/update')
async def material_price_update(
    request: Request,
    payload: Annotated[dict, Body()],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    await _ensure_tables(db)
    item_id = _to_int(payload.get('id'), 0)
    if item_id <= 0:
        return ResponseUtil.failure(msg='ID不能为空')

    current_rows = await _fetch_all(
        db,
        """
        SELECT id,
               material_name,
               specification,
               category_level3_id,
               process_segment_name,
               unit_name,
               remark,
               standardization_status,
               standardization_error
        FROM material_price_record
        WHERE id = :id AND del_flag = '0'
        LIMIT 1
        """,
        {'id': item_id},
    )
    if not current_rows:
        return ResponseUtil.failure(msg='材价记录不存在')
    current_row = current_rows[0]

    normalized = _normalize_row(payload)
    check = await _validate_rows(db, [normalized], with_db_duplicate=False)
    row_check = check['rows'][0]
    if row_check['errors']:
        return ResponseUtil.failure(msg='校验不通过', data=row_check)

    if await _exists_duplicate_in_db(db, normalized, exclude_id=item_id):
        return ResponseUtil.failure(msg='数据库已存在同一业务键记录')

    await _upsert_import_dicts(db, [normalized])

    changed_fields = _get_restandardize_changed_fields(current_row, normalized)
    should_restandardize = len(changed_fields) > 0
    target_status = (
        STANDARD_STATUS_PENDING
        if should_restandardize
        else _as_str(current_row.get('standardization_status')) or STANDARD_STATUS_PENDING
    )

    update_by = _as_str(payload.get('updateBy')) or 'system'
    await db.execute(
        text(
            """
            UPDATE material_price_record
            SET material_name = :material_name,
                specification = :specification,
                category_name = :category_name,
                category_level1_id = :category_level1_id,
                category_level2_id = :category_level2_id,
                category_level3_id = :category_level3_id,
                process_segment_name = :process_segment_name,
                unit_name = :unit_name,
                brand_name = :brand_name,
                quantity = :quantity,
                price_excluding_tax = :price_excluding_tax,
                tax_rate = :tax_rate,
                price_including_tax = :price_including_tax,
                source_project = :source_project,
                purchase_time = :purchase_time,
                price_type = :price_type,
                supplier_company = :supplier_company,
                remark = :remark,
                update_by = :update_by,
                standardization_status = :standardization_status,
                standardization_error = CASE
                    WHEN :should_restandardize = 1 THEN NULL
                    ELSE standardization_error
                END
            WHERE id = :id
            """
        ),
        {
            'id': item_id,
            'material_name': normalized['materialName'],
            'specification': normalized['specification'],
            'category_name': normalized['categoryName'],
            'category_level1_id': normalized.get('categoryLevel1Id') or None,
            'category_level2_id': normalized.get('categoryLevel2Id') or None,
            'category_level3_id': normalized.get('categoryLevel3Id') or None,
            'process_segment_name': normalized['processSegmentName'] or None,
            'unit_name': normalized['unit'],
            'brand_name': normalized['brand'] or None,
            'quantity': str(_to_decimal(normalized['quantity'], Decimal('0'))),
            'price_excluding_tax': str(_round2(_to_decimal(normalized['priceExcludingTax'], Decimal('0')))),
            'tax_rate': str(_round2(_to_decimal(normalized['taxRate'], Decimal('0')))),
            'price_including_tax': str(_round2(_to_decimal(normalized['priceIncludingTax'], Decimal('0')))),
            'source_project': normalized['sourceProject'] or None,
            'purchase_time': normalized['purchaseTime'],
            'price_type': normalized['priceType'],
            'supplier_company': normalized['supplierCompany'] or None,
            'remark': normalized['remark'] or None,
            'update_by': update_by,
            'standardization_status': target_status,
            'should_restandardize': 1 if should_restandardize else 0,
        },
    )

    if should_restandardize:
        # 二次触发标准化时，同步刷新标准化复核中的原始材料信息，再入队Dify
        await create_review_task_for_price(db, item_id)

    await db.commit()

    if should_restandardize:
        enqueue_standardization([item_id])

    return ResponseUtil.success(
        msg='更新成功',
        data={
            'restandardized': should_restandardize,
            'changedFields': changed_fields,
        },
    )


@material_price_controller.post('/delete')
async def material_price_delete(
    request: Request,
    payload: Annotated[dict, Body()],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    await _ensure_tables(db)
    item_id = _to_int(payload.get('id'), 0)
    if item_id <= 0:
        return ResponseUtil.failure(msg='ID不能为空')

    await db.execute(
        text("UPDATE material_price_record SET del_flag = '1' WHERE id = :id"),
        {'id': item_id},
    )
    await db.commit()
    return ResponseUtil.success(msg='删除成功')
